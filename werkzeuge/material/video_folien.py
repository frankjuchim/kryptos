"""Erklärvideo-Folien: materialien/Videos/NN_Kurz/video.json → Präsentationsquelle → .pptx (Kamera-Zone frei).

    python3 werkzeuge/material/video_folien.py [NN_Kurz …]      # ohne Argument: alle Videos
    python3 werkzeuge/material/video_folien.py --pruefen [NN_Kurz …]
        baut in einen Temp-Ordner und meldet (Exit 1): Inhalt in der Kamera-Zone (außer Titelfolie), Bildplatzhalter,
        Kontrast-Verstöße (kr_kontrast), Folientext/Notizen ≠ Repo-Datei („nicht aktuell“), Quelle
        quellen/material/Videos/video_NN.json nicht aktuell, Sprechertext nicht wörtlich in Skript.md (Plan M9 Task 214)

Quellen werden nach quellen/material/Videos/video_NN.json geschrieben und mit bauen.py gebaut; die Folien liegen danach
neben der video.json (materialien/Videos/NN_Kurz/KR_Video_NN_Kurz_Folien.pptx). Sprechertext und Regie stehen in den
Sprechernotizen.
"""
import glob
import json
import os
import re
import sys

HIER = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HIER, "..", ".."))
sys.path.insert(0, HIER)
import bauen  # noqa: E402


def _ohne_md(t):
    return re.sub(r"\*\*|`|==", "", str(t))


def kasten(k):
    if not k:
        return None
    text = k.get("text", [])
    if isinstance(text, str):
        text = [text]
    return {"t": "hinweis", "kopf": k.get("titel", "").upper(), "farbe": k.get("farbe", "hinweis"),
            "inhalt": [{"t": "p", "md": z} for z in text]}


def punkte_bloecke(punkte):
    return [{"t": "liste", "items": punkte}] if punkte else []


def mit_kasten(haupt, k, teilung=None):
    kb = kasten(k)
    if not kb:
        return haupt
    anteil = 0.62 if not teilung else max(0.45, min(0.72, float(teilung) / 10))
    return [{"t": "zweispaltig", "anteil": anteil, "links": haupt, "rechts": [kb]}]


def bildblock(name, v):
    b = (v.get("bilder") or {}).get(name) if isinstance(name, str) else name
    if not b:
        return {"t": "platzhalter", "beschreibung": str(name)}
    t = b.get("t") or b.get("typ")
    if t == "haeufigkeit":
        return {"t": "haeufigkeit", "werte": b.get("werte"), "text": b.get("text"), "titel": b.get("titel"),
                "sortiert": b.get("sortiert", False), "referenz": bool(b.get("referenz")), "hoehe": 3.2}
    if t == "quadrat":
        hervor = []
        m = b.get("markiere") or {}
        for z in m.get("zeilen", []):
            r = ord(z) - 65
            hervor += [[r, s] for s in range(26)]
        for sp in m.get("spalten", []):
            c = ord(sp) - 65
            hervor += [[r, c] for r in range(26)]
        return {"t": "quadrat", "art": b.get("art", "vigenere"), "hervor": hervor} if b.get("art", "vigenere") == "vigenere" \
            else {"t": "quadrat", "zeilen": b.get("zeilen"), "kopfzeile": b.get("kopfzeile"), "kopfspalte": b.get("kopfspalte")}
    if t == "bild":
        # Bilddatei relativ zum Videoordner (Plan M9 Task 213, E156) → Pfad relativ zum Repo für bauen.py
        ordner = f"{int(v['nr']):02d}_{v['kurz']}"
        # Objektbilder des Kryptolabors (Plan M10 Task 240) liegen zentral unter materialien/Bilder/Labor/
        pfad = b["pfad"] if b["pfad"].startswith("materialien/") else f"materialien/Videos/{ordner}/{b['pfad']}"
        block = {"t": "bild", "pfad": pfad}
        if b.get("beschreibung"):
            block["alt"] = b["beschreibung"]
        if b.get("unterschrift"):
            block["unterschrift"] = b["unterschrift"]
        return block
    return {"t": "platzhalter", "beschreibung": b.get("beschreibung", name)}


def tabelle(tb):
    t = {"t": "tabelle", "kopf": tb.get("kopf"), "zeilen": tb["zeilen"], "breiten": tb.get("breiten"),
         "mono": tb.get("mono", False), "zentriert": tb.get("zentriert", False)}
    if tb.get("groesse"):
        t["schrift_folie"] = tb["groesse"]
    if tb.get("hervorheben"):
        t["hervor"] = [[r, c] for r in tb["hervorheben"] for c in range(len(tb["zeilen"][0]))]
    return t


def schaubild(sb):
    ids = {}
    for a in sb["akteure"]:
        ids[a["id"]] = a["id"]
        ids[a["name"]] = a["id"]
    schritte = []
    for p in sb.get("pfeile", sb.get("schritte", [])):
        von, nach = ids.get(p["von"]), ids.get(p["nach"])
        if von is None:
            von = nach
        if nach is None:
            nach = von
        schritte.append({"nr": p.get("nr", ""), "von": von, "nach": nach, "text": _ohne_md(p["text"])})
    akteure = [{"id": a["id"], "name": a["name"], "unter": a.get("rolle")} for a in sb["akteure"]]
    return {"t": "schaubild", "akteure": akteure, "schritte": schritte, "titel": sb.get("titel")}


def notizen(f):
    teile = []
    if f.get("kapitel"):
        teile.append(f"[Kapitel: {f['kapitel']}]")
    if f.get("regie"):
        teile.append(f"REGIE: {f['regie']}")
    if f.get("sprechertext"):
        teile.append(f"SPRECHERTEXT: {f['sprechertext']}")
    if f.get("pause_s"):
        teile.append(f"PAUSE: {f['pause_s']} s Einblendung „⏸ Video anhalten“")
    return "\n\n".join(teile)


def folie(f, v):
    typ = f["typ"]
    basis = {"titel": _ohne_md(f.get("titel", "")), "notizen": notizen(f), "video": True,
             "abschnitt": f.get("kapitel")}
    if typ == "titel":
        return {**basis, "layout": "titel", "titel": v["titel"], "untertitel": v.get("untertitel", ""),
                "oberzeile": f"VIDEO {int(v['nr']):02d} · " + f.get("ort", "").upper(), "video": False}
    if typ == "punkte":
        haupt = punkte_bloecke(f.get("punkte", []))
        if f.get("bild"):
            bb = bildblock(f["bild"], v)
            if bb["t"] == "platzhalter":
                bb["hoehe"] = 1.6
            elif bb["t"] == "bild":
                bb["_maxH"] = 1.9  # unter den Stichpunkten (Folie bleibt unter 5,3 Zoll Inhalt)
            haupt.append(bb)
        return {**basis, "layout": "inhalt", "inhalt": mit_kasten(haupt, f.get("kasten"), f.get("teilung"))}
    if typ == "bild":
        haupt = [bildblock(f["bild"], v)]
        if haupt[0]["t"] == "bild" and haupt[0].get("unterschrift"):
            # Objektbild im Seitenverhältnis des Objekts (M10-Review Fund 7) + Unterschrift bleibt unter 5,3 Zoll Inhalt
            haupt[0]["_maxH"] = 3.9 if f.get("text") else 4.6
        if f.get("text"):
            haupt.insert(0, {"t": "p", "md": f["text"]})
            if haupt[1]["t"] == "quadrat":
                haupt[1]["max_hoehe"] = 4.3
        return {**basis, "layout": "inhalt", "inhalt": mit_kasten(haupt, f.get("kasten"), f.get("teilung"))}
    if typ == "tabelle":
        haupt = []
        if f.get("text"):
            haupt.append({"t": "p", "md": f["text"]})
        haupt.append(tabelle(f["tabelle"]))
        for u in f.get("unten", []) if isinstance(f.get("unten"), list) else ([f["unten"]] if f.get("unten") else []):
            haupt.append({"t": "p", "md": u})
        lang = len(f["tabelle"]["zeilen"]) >= 10
        if lang:
            haupt[-1 if not f.get("unten") else 1]["schrift_folie"] = min(f["tabelle"].get("groesse", 14), 13)
        return {**basis, "layout": "inhalt",
                "inhalt": mit_kasten(haupt, f.get("kasten"), 7.0 if lang else f.get("teilung"))}
    if typ == "geheimtext":
        haupt = []
        if f.get("text"):
            haupt.append({"t": "p", "md": f["text"]})
        haupt.append({"t": "geheimtext", "text": f["geheimtext"], "groesse": f.get("groesse", 20)})
        if f.get("zwischen"):
            haupt.append({"t": "p", "md": f["zwischen"]})
        if f.get("klartext"):
            haupt.append({"t": "klartext", "text": f["klartext"]})
        return {**basis, "layout": "inhalt", "inhalt": mit_kasten(haupt, f.get("kasten"))}
    if typ == "alphabet":
        a = f["alphabet"]
        zeilen = []
        if a.get("geheim"):
            zeilen.append({"label": a.get("geheim_titel", "Geheim"), "werte": a["geheim"]})
        haupt = [{"t": "alphabet", "nummerierung": 1 if "1" in str(a.get("nummerierung", "A = 1")) else 0,
                  "zeilen": zeilen}]
        for u in f.get("unten", []):
            haupt.append({"t": "p", "md": u})
        return {**basis, "layout": "inhalt", "inhalt": haupt}
    if typ == "schaubild":
        return {**basis, "layout": "inhalt", "inhalt": mit_kasten([schaubild(f["schaubild"])], f.get("kasten"))}
    if typ == "falle":
        fa, ri = f.get("falsch", {}), f.get("richtig", {})
        return {**basis, "layout": "falle", "nr": f["falle"], "titel": f.get("kurztitel", ""), "regel": f.get("regel"),
                **({"schrift": f["schrift"]} if f.get("schrift") else {}),
                "falsch": [{"t": "p", "md": z} for z in fa.get("zeilen", [])],
                "richtig": [{"t": "p", "md": z} for z in ri.get("zeilen", [])]}
    if typ == "pause":
        haupt = [{"t": "p", "md": f.get("aufgabe", "")}]
        if f.get("geheimtext"):
            haupt.append({"t": "geheimtext", "text": f["geheimtext"]})
        if f.get("klartext"):
            haupt.append({"t": "klartext", "text": f["klartext"]})
        if f.get("tabelle"):
            haupt.append(tabelle(f["tabelle"]))
        if f.get("code"):
            haupt.append({"t": "code", "text": f["code"]})
        return {**basis, "layout": "auftrag", "sozialform": "⏸ Video anhalten", "zeit": "selbst probieren",
                "inhalt": mit_kasten(haupt, f.get("kasten"))}
    if typ == "merke":
        return {**basis, "layout": "merke", "inhalt": [{"t": "p", "md": z} for z in f.get("punkte", [])]}
    if typ == "ende":
        links = f.get("links", {})
        haupt = [{"t": "tabelle", "zeilen": [[k, v2] for k, v2 in links.items()], "kopfspalte": True,
                  "breiten": [2.0, 6.0]}]
        if f.get("naechstes"):
            haupt.append({"t": "p", "md": f"**Nächstes Video:** {f['naechstes']}"})
        return {**basis, "layout": "inhalt", "titel": "Weiter im Lab", "inhalt": haupt}
    raise ValueError("unbekannter Folientyp " + typ)


def quelle(ordner):
    """Präsentationsquelle eines Videos (ohne zu schreiben)."""
    v = json.load(open(os.path.join(ordner, "video.json"), encoding="utf-8"))
    kurz = os.path.basename(ordner)
    return {"typ": "praesentation", "welt": v.get("welt", 9) if isinstance(v.get("welt"), int) else 9,
            "titel": v["titel"], "datei": f"KR_Video_{kurz}_Folien",
            "ordner": os.path.relpath(ordner, os.path.join(REPO, "materialien")),
            "folien": [folie(f, v) for f in v["folien"]]}


def quellpfad(ordner):
    return os.path.join(REPO, "quellen", "material", "Videos", f"video_{os.path.basename(ordner)[:2]}.json")


def konvertiere(ordner):
    q = quelle(ordner)
    ziel = quellpfad(ordner)
    os.makedirs(os.path.dirname(ziel), exist_ok=True)
    json.dump(q, open(ziel, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return ziel


# ---------------------------------------------------------------------------------------------------
# Prüfung (Plan M9 Task 214)
# ---------------------------------------------------------------------------------------------------

EMU_ZOLL = 914400
KAMERA_X, KAMERA_Y = 8.61, 4.74  # Kamera-Zone unten rechts (LIESMICH): ab x = 8,61 in und y = 4,74 in


def folientexte(pfad):
    """Je Folie: (Texte der Formen in Reihenfolge, Notizen) – für den Vergleich alt ↔ neu."""
    from pptx import Presentation
    aus = []
    for folie_ in Presentation(pfad).slides:
        texte = []
        for sh in folie_.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texte.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                texte.append(" | ".join(c.text for r in sh.table.rows for c in r.cells))
        notiz = folie_.notes_slide.notes_text_frame.text.strip() if folie_.has_notes_slide else ""
        aus.append((texte, notiz))
    return aus


def pruefe_pptx_datei(pfad):
    """Befunde: Bildplatzhalter, Inhalt in der Kamera-Zone (Folien nach der Titelfolie)."""
    from pptx import Presentation
    befunde = []
    for nr, folie_ in enumerate(Presentation(pfad).slides, 1):
        for sh in folie_.shapes:
            # Foliennummer des Masters liegt unter dem Kamerabild – kein Inhalt (wird in der Aufnahme verdeckt)
            if sh.is_placeholder and sh.placeholder_format.type is not None and str(sh.placeholder_format.type).startswith("SLIDE_NUMBER"):
                continue
            if sh.name == "Bildplatzhalter":
                befunde.append(f"Folie {nr}: Bildplatzhalter")
            if nr == 1 or sh.left is None or sh.width is None:
                continue
            rechts = (sh.left + sh.width) / EMU_ZOLL
            unten = (sh.top + sh.height) / EMU_ZOLL
            if rechts > KAMERA_X + 0.005 and unten > KAMERA_Y + 0.005:
                befunde.append(f"Folie {nr}: „{sh.name}“ ragt in die Kamera-Zone (rechts {rechts:.2f} in, unten {unten:.2f} in)")
    return befunde


def _norm(t):
    return re.sub(r"\s+", " ", t).strip()


def pruefen(ordner_liste):
    import tempfile
    import kr_kontrast
    alle = 0
    for o in ordner_liste:
        name = os.path.basename(o)
        befunde = []
        v = json.load(open(os.path.join(o, "video.json"), encoding="utf-8"))
        q = quelle(o)
        # Quelle aktuell?
        qp = quellpfad(o)
        if not os.path.exists(qp) or json.load(open(qp, encoding="utf-8")) != json.loads(json.dumps(q, ensure_ascii=False)):
            befunde.append(f"{os.path.relpath(qp, REPO)} nicht aktuell (video_folien.py {name})")
        with tempfile.TemporaryDirectory() as tmp:
            tq = os.path.join(tmp, "quelle.json")
            json.dump(q, open(tq, "w", encoding="utf-8"), ensure_ascii=False)
            dateien, warn = bauen.baue(tq, os.path.join(tmp, "materialien"))
            neu = dateien[0]
            befunde += [f"Generator: {w}" for w in warn]
            befunde += pruefe_pptx_datei(neu)
            befunde += [f"Kontrast: {b}" for b in kr_kontrast.pruefe(neu)]
            repo = os.path.join(o, os.path.basename(neu))
            if not os.path.exists(repo):
                befunde.append(f"{os.path.relpath(repo, REPO)} fehlt")
            elif folientexte(repo) != folientexte(neu):
                befunde.append(f"{os.path.relpath(repo, REPO)} nicht aktuell (Folientext oder Notizen weichen ab)")
        # Sprechertexte wörtlich im Drehbuch
        skript_pfad = os.path.join(o, "Skript.md")
        skript = _norm(open(skript_pfad, encoding="utf-8").read()) if os.path.exists(skript_pfad) else ""
        for i, f in enumerate(v["folien"], 1):
            if f.get("sprechertext") and _norm(f["sprechertext"]) not in skript:
                befunde.append(f"Folie {i}: Sprechertext steht nicht wörtlich in Skript.md")
        for b in befunde:
            print(f"✘ {name}: {b}")
        if not befunde:
            print(f"✔ {name}")
        alle += len(befunde)
    print(f"{len(ordner_liste)} Videos, {alle} Befunde")
    return 1 if alle else 0


if __name__ == "__main__":
    argumente = [a for a in sys.argv[1:] if a != "--pruefen"]
    ordner = argumente or sorted(d for d in glob.glob(os.path.join(REPO, "materialien", "Videos", "[0-9][0-9]_*")) if os.path.isdir(d))
    ordner = [o if os.path.isabs(o) else os.path.join(REPO, "materialien", "Videos", o) for o in ordner]
    if "--pruefen" in sys.argv[1:]:
        sys.exit(pruefen(ordner))
    for o in ordner:
        q = konvertiere(o)
        dateien, warn = bauen.baue(q, os.path.join(REPO, "materialien"))
        print("✔", os.path.relpath(dateien[0], REPO))
        for w in warn:
            print("  ⚠", w)
